from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sample_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Название кампании", "Название группы", "Заголовок 1", "Текст", "Ключевые фразы"])
    ws.append([
        "Кампания 1", "Группа A", "Купить iPhone 15",
        "Закажите сегодня — доставка 1 день, гарантия 2 года.",
        "iphone 15, купить iphone",
    ])
    ws.append([
        "", "", "Samsung S24 со скидкой",
        "Доставка за 2 часа по Москве. Оформите заказ онлайн.",
        "samsung s24",
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------
# /excel/upload
# ---------------------------------------------------------------------
def test_upload_requires_auth(anon_client):
    files = {"file": ("ads.xlsx", _sample_xlsx(), _XLSX_MIME)}
    response = anon_client.post("/excel/upload", files=files)
    assert response.status_code == 401


def test_upload_happy_path(auth_client, monkeypatch):
    # Mock AI improvement so the test doesn't need an OpenAI key.
    async def _fake_improve(ad):
        return {
            "headline": "Улучшенный заголовок",
            "text": "Улучшенный текст с призывом к действию.",
            "reasoning": "Добавлен CTA и конкретика.",
        }

    monkeypatch.setattr("app.services.ai_ads.improve_ad", _fake_improve)

    files = {"file": ("ads.xlsx", _sample_xlsx(), _XLSX_MIME)}
    response = auth_client.post("/excel/upload", files=files)
    assert response.status_code == 200

    body = response.json()
    assert body["summary"]["total_ads"] == 2
    assert body["summary"]["total_campaigns"] == 1
    assert body["summary"]["improved_count"] == 2
    assert body["insights"]["estimated_ctr_loss"] in {
        "low", "moderate (~10-20%)", "high (~20-35%)", "severe (>35%)", "n/a",
    }

    first = body["ads"][0]
    assert first["original"]["headline"] == "Купить iPhone 15"
    assert 0 <= first["audit"]["score"] <= 100
    assert first["improved"]["headline"] == "Улучшенный заголовок"


def test_upload_rejects_non_xlsx(auth_client):
    files = {"file": ("not-an-excel.txt", b"hello", "text/plain")}
    response = auth_client.post("/excel/upload", files=files)
    assert response.status_code == 400
    assert response.json()["detail"] == "unsupported_file_type"


def test_upload_broken_xlsx_returns_errors_not_500(auth_client):
    files = {"file": ("broken.xlsx", b"not really xlsx bytes", _XLSX_MIME)}
    response = auth_client.post("/excel/upload", files=files)
    assert response.status_code == 200

    body = response.json()
    assert body["summary"]["total_ads"] == 0
    assert body["errors"]
    assert body["errors"][0]["field"] == "file"


def test_upload_ai_failure_does_not_fail_whole_response(auth_client, monkeypatch):
    # improve_ad raises → surfaced as `improved: null`, rest intact.
    async def _broken(_ad):
        raise RuntimeError("upstream boom")

    monkeypatch.setattr("app.services.ai_ads.improve_ad", _broken)

    files = {"file": ("ads.xlsx", _sample_xlsx(), _XLSX_MIME)}
    response = auth_client.post("/excel/upload", files=files)
    assert response.status_code == 200

    body = response.json()
    assert body["summary"]["improved_count"] == 0
    assert all(ad["improved"] is None for ad in body["ads"])


# ---------------------------------------------------------------------
# /excel/export
# ---------------------------------------------------------------------
def test_export_requires_auth(anon_client):
    payload = {"ads": [{"headline": "h", "text": "t"}]}
    response = anon_client.post("/excel/export", json=payload)
    assert response.status_code == 401


def test_export_round_trip(auth_client):
    payload = {
        "filename": "Моя_Кампания",
        "ads": [
            {
                "campaign": "Кампания 1",
                "group": "Группа A",
                "headline": "Купить iPhone 15",
                "headline2": "Скидка 20%",
                "text": "Закажите сегодня — доставка 1 день, гарантия 2 года.",
                "keywords": ["iphone 15", "купить iphone"],
            }
        ],
    }
    response = auth_client.post("/excel/export", json=payload)
    assert response.status_code == 200
    assert response.headers["content-type"] == _XLSX_MIME
    assert "attachment" in response.headers["content-disposition"]
    assert "Моя_Кампания.xlsx" in response.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    assert "Объявления" in wb.sheetnames
    assert "Кампании" in wb.sheetnames


def test_export_rejects_empty_ads(auth_client):
    response = auth_client.post("/excel/export", json={"ads": []})
    assert response.status_code == 422  # Pydantic min_length=1
