from __future__ import annotations

from app.services.excel_export import build_direct_xlsx, slugify_filename
from app.services.excel_import import _parse_workbook_bytes


def test_round_trip_export_import():
    ads = [
        {
            "campaign": "Кампания 1",
            "group": "Группа A",
            "headline": "Купить iPhone 15",
            "headline2": "Скидка 20%",
            "text": "Закажите сегодня — доставка 1 день, гарантия 2 года.",
            "keywords": ["iphone 15", "купить iphone"],
        },
        {
            "campaign": "Кампания 1",
            "group": "Группа A",
            "headline": "Samsung S24 — заказ онлайн",
            "headline2": None,
            "text": "Доставка за 2 часа по Москве.",
            "keywords": ["samsung s24"],
        },
    ]

    data = build_direct_xlsx(ads)
    assert data[:2] == b"PK"  # xlsx is a zip

    parsed = _parse_workbook_bytes(data)
    assert parsed["errors"] == []
    assert len(parsed["ads"]) == 2
    assert parsed["ads"][0]["headline"] == "Купить iPhone 15"
    assert parsed["ads"][0]["campaign"] == "Кампания 1"
    assert "iphone 15" in parsed["ads"][0]["keywords"]


def test_trims_overlong_fields():
    ads = [{
        "campaign": "",
        "group": "",
        "headline": "Слово " * 30,  # far > 56
        "text": "Тоже слово " * 30,  # far > 81
        "keywords": [],
    }]
    data = build_direct_xlsx(ads)
    parsed = _parse_workbook_bytes(data)

    assert parsed["errors"] == []  # trimmed to within limits
    assert len(parsed["ads"][0]["headline"]) <= 56
    assert len(parsed["ads"][0]["text"]) <= 81


def test_slugify_filename():
    assert slugify_filename(None) == "direct_ads"
    assert slugify_filename("") == "direct_ads"
    assert slugify_filename("My Campaign / v2") == "My_Campaign_v2"
    # Cyrillic is preserved; dangerous characters collapsed.
    assert slugify_filename("Кампания №1!") == "Кампания_1"


def test_campaigns_sheet_is_present():
    ads = [
        {"campaign": "K1", "group": "G1", "headline": "h", "text": "t", "keywords": []},
        {"campaign": "K1", "group": "G2", "headline": "h", "text": "t", "keywords": []},
        {"campaign": "K2", "group": "",   "headline": "h", "text": "t", "keywords": []},
    ]
    data = build_direct_xlsx(ads)
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    assert "Кампании" in wb.sheetnames
    sheet = wb["Кампании"]
    rows = list(sheet.iter_rows(values_only=True))
    # Header + two campaigns (K1, K2).
    assert len(rows) == 3
    # Headers are styled (at least the text matches).
    assert rows[0][0] == "Название кампании"
