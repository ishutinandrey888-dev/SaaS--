"""Generate a Yandex Direct-shaped XLSX from a list of ads.

Companion to `excel_import`: the pair lets a user upload → audit →
rewrite → download, and import the result back into Yandex Direct.

The output is deliberately minimal — just the columns that identify an
ad — so it round-trips cleanly through our own parser and is easy to
paste into the Direct XLS Master template.
"""

from __future__ import annotations

import io
import re
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.excel_import import HEADLINE_MAX, HEADLINE2_MAX, TEXT_MAX

_HEADER_FILL = PatternFill("solid", fgColor="E8E8E8")
_HEADER_FONT = Font(bold=True)
_HEADER_ALIGN = Alignment(vertical="center", wrap_text=True)
_CELL_ALIGN = Alignment(vertical="top", wrap_text=True)

_ADS_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Название кампании", 28),
    ("Название группы",   28),
    ("Заголовок 1",       45),
    ("Заголовок 2",       28),
    ("Текст объявления",  55),
    ("Ключевые фразы",    45),
)

_CAMPAIGNS_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Название кампании", 32),
    ("Группы объявлений", 50),
    ("Количество объявлений", 18),
)

_FILENAME_SAFE_RE = re.compile(r"[^A-Za-z0-9а-яА-ЯёЁ_\-]+")


def slugify_filename(name: str | None, *, default: str = "direct_ads") -> str:
    """Return a safe XLSX filename (without extension)."""
    if not name:
        return default
    cleaned = _FILENAME_SAFE_RE.sub("_", name.strip()).strip("_")
    return cleaned[:80] or default


def _trim(value: str | None, limit: int) -> str:
    if not value:
        return ""
    text = str(value).strip()
    if len(text) <= limit:
        return text
    snippet = text[:limit]
    if " " in snippet:
        snippet = snippet.rsplit(" ", 1)[0]
    return snippet.strip()


def _style_header(sheet, column_count: int) -> None:
    for col in range(1, column_count + 1):
        cell = sheet.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"


def _write_ads_sheet(sheet, ads: Iterable[dict[str, Any]]) -> None:
    sheet.title = "Объявления"
    headers = [h for h, _ in _ADS_COLUMNS]
    sheet.append(headers)

    for idx, (_, width) in enumerate(_ADS_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    for ad in ads:
        keywords = ad.get("keywords") or []
        if isinstance(keywords, str):
            keywords_text = keywords
        else:
            keywords_text = "\n".join(str(k).strip() for k in keywords if str(k).strip())

        row_values = [
            str(ad.get("campaign") or "").strip(),
            str(ad.get("group") or "").strip(),
            _trim(ad.get("headline"), HEADLINE_MAX),
            _trim(ad.get("headline2"), HEADLINE2_MAX),
            _trim(ad.get("text"), TEXT_MAX),
            keywords_text,
        ]
        sheet.append(row_values)

    _style_header(sheet, len(_ADS_COLUMNS))

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _CELL_ALIGN


def _write_campaigns_sheet(sheet, ads: list[dict[str, Any]]) -> None:
    sheet.title = "Кампании"
    headers = [h for h, _ in _CAMPAIGNS_COLUMNS]
    sheet.append(headers)

    for idx, (_, width) in enumerate(_CAMPAIGNS_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    buckets: dict[str, dict[str, Any]] = {}
    for ad in ads:
        name = str(ad.get("campaign") or "").strip() or "(без названия)"
        entry = buckets.setdefault(name, {"groups": set(), "count": 0})
        group = str(ad.get("group") or "").strip()
        if group:
            entry["groups"].add(group)
        entry["count"] += 1

    for name, entry in buckets.items():
        groups_text = ", ".join(sorted(entry["groups"]))
        sheet.append([name, groups_text, entry["count"]])

    _style_header(sheet, len(_CAMPAIGNS_COLUMNS))

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _CELL_ALIGN


def build_direct_xlsx(ads: list[dict[str, Any]]) -> bytes:
    """Return XLSX bytes with two sheets: Объявления + Кампании."""
    wb = Workbook()
    _write_ads_sheet(wb.active, ads)
    _write_campaigns_sheet(wb.create_sheet(), ads)

    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
